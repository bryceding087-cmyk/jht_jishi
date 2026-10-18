"""Validate uploaded storyboard workbooks without storing or changing projects."""

from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata
from zipfile import ZipFile
from xml.etree import ElementTree

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ENTRY_BYTES = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 1024
MAX_SHEET_ROWS = 10000
MAX_SHEET_COLUMNS = 128
MAX_SCENES = 80
MAX_CELL_LENGTH = 4000
HEADERS = ('镜号', '时长', '画面描述', '景别', '光影氛围', '对白·旁白', '音效', '运镜', '最终提示词')
FIELDS = ('index', 'duration', 'description', 'shot_size', 'lighting', 'dialogue', 'sound', 'camera', 'final_prompt')


def invalid(message, status=422):
    raise HTTPException(status, message)


def normalise_header(value):
    text = unicodedata.normalize('NFKC', str(value or '')).strip()
    text = re.sub(r'\s+', '', text)
    if re.fullmatch(r'对白[·•・/／、.\-—与及]*旁白', text):
        return '对白·旁白'
    return text


def _check_archive(data):
    if data.startswith(b'\xd0\xcf\x11\xe0'):
        invalid('不支持加密文件或旧版 Excel，请取消密码并另存为 .xlsx 后重试')
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                invalid('Excel 文件内部条目过多，请使用下载的模板重新整理', 413)
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)):
                invalid('Excel 文件结构异常，包含重复文件，请重新另存为 .xlsx')
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                invalid('文件不是有效的 Excel .xlsx 工作簿')
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                invalid('Excel 解压后超过 50 MB，请减小文件后重试', 413)
            for entry in entries:
                if entry.flag_bits & 1:
                    invalid('不支持加密 Excel，请取消密码后重试')
                if entry.file_size > MAX_ENTRY_BYTES:
                    invalid('Excel 内部单个文件超过 20 MB，请减小文件后重试', 413)
                with archive.open(entry) as stream:
                    content = stream.read(MAX_ENTRY_BYTES + 1)
                if len(content) > MAX_ENTRY_BYTES:
                    invalid('Excel 内部单个文件超过 20 MB，请减小文件后重试', 413)
                if entry.filename.endswith('.xml'):
                    # Reject entity declarations before the spreadsheet reader parses XML.
                    if re.search(br'<!\s*(?:DOCTYPE|ENTITY)\b', content.replace(b'\x00', b''), re.IGNORECASE):
                        invalid('Excel 包含不支持的 XML 内容，请重新另存为 .xlsx')
                if re.fullmatch(r'xl/worksheets/[^/]+\.xml', entry.filename):
                    for _, element in ElementTree.iterparse(BytesIO(content), events=('end',)):
                        if element.tag.endswith('}row'):
                            row = int(element.get('r', '0'))
                            if row < 1 or row > MAX_SHEET_ROWS:
                                invalid('工作表行范围过大，请移除第 10000 行之后的内容或使用模板', 413)
                            element.clear()
                        elif element.tag.endswith('}c'):
                            coordinate = element.get('r', '')
                            match = re.fullmatch(r'([A-Z]+)(\d+)', coordinate)
                            if not match:
                                invalid('Excel 单元格坐标异常，请重新另存为 .xlsx')
                            column = 0
                            for character in match[1]:
                                column = column * 26 + ord(character) - 64
                            if column > MAX_SHEET_COLUMNS or int(match[2]) > MAX_SHEET_ROWS:
                                invalid('工作表范围过大，请只保留分镜表格所需的行列', 413)
                            element.clear()
    except HTTPException:
        raise
    except Exception:
        invalid('Excel 文件已损坏或无法读取，请重新另存为 .xlsx 后上传')


def _cell_text(cell, row, column, label):
    location = f'第 {row} 行「{label}」列（{get_column_letter(column)}{row}）'
    if cell.data_type == 'f':
        invalid(f'{location}含公式，请粘贴为纯文本或数值后导入')
    if cell.data_type == 'e':
        invalid(f'{location}含 Excel 错误值，请修正后导入')
    if isinstance(cell.value, (date, datetime, time)):
        invalid(f'{location}为日期或时间格式，请改为纯文本；时长请填写秒数')
    if isinstance(cell.value, bool):
        invalid(f'{location}为逻辑值，请改为纯文本或数值')
    text = '' if cell.value is None else str(cell.value).strip()
    if len(text) > MAX_CELL_LENGTH:
        invalid(f'{location}超过 4000 字，请缩短内容')
    return text


def _positive_decimal(text, row, label, integer=False):
    cleaned = unicodedata.normalize('NFKC', text).strip()
    if not integer:
        cleaned = re.sub(r'\s*(?:s|秒)\s*$', '', cleaned, flags=re.IGNORECASE)
    try:
        value = Decimal(cleaned)
        if not value.is_finite() or value <= 0 or (integer and value != value.to_integral_value()):
            raise InvalidOperation
        # Bound expanded notation to avoid allocating enormous strings from 1e999999.
        if value.adjusted() >= MAX_CELL_LENGTH or value.as_tuple().exponent < -(MAX_CELL_LENGTH - 2):
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        hint = '正整数' if integer else '大于 0 的秒数，例如 5、5s 或 5秒'
        invalid(f'第 {row} 行「{label}」请填写{hint}')
    return value


def parse_storyboard_excel(data, file_name):
    """Return an import preview; callers decide when to apply it to a node."""
    if len(data) > MAX_FILE_BYTES:
        invalid('Excel 文件不能超过 5 MB', 413)
    if not file_name or not str(file_name).lower().endswith('.xlsx'):
        invalid('仅支持 .xlsx 文件，请下载模板填写或将文件另存为 .xlsx')
    if not data:
        invalid('上传的 Excel 文件为空')
    _check_archive(data)
    workbook = None
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        if not workbook.worksheets:
            invalid('Excel 中没有可读取的工作表')
        sheet = workbook['分镜脚本'] if '分镜脚本' in workbook.sheetnames else workbook.worksheets[0]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_col=MAX_SHEET_COLUMNS)
        header_row = next(rows, ())
        columns = {}
        extra_headers = []
        for column, cell in enumerate(header_row, 1):
            name = normalise_header(_cell_text(cell, 1, column, '表头'))
            if not name:
                continue
            if name in columns:
                invalid(f'第 1 行存在重复表头「{name}」，请保留一列')
            columns[name] = column
            if name not in HEADERS:
                extra_headers.append(name)
        missing = [name for name in HEADERS if name not in columns]
        if missing:
            invalid(f'第 1 行缺少表头：{"、".join(missing)}。请按下载模板填写')
        scenes = []
        used_indices = set()
        renumbered = False
        for row_number, row in enumerate(rows, 2):
            # Ignore genuinely empty rows, including styled cells in the template.
            if not any(cell.value is not None and str(cell.value).strip() for cell in row):
                continue
            if len(scenes) >= MAX_SCENES:
                invalid(f'分镜数量超过 80 条（第 {row_number} 行），请拆分文件后导入')
            values = {
                field: _cell_text(row[columns[header] - 1], row_number, columns[header], header)
                for header, field in zip(HEADERS, FIELDS)
            }
            original_index = int(_positive_decimal(values['index'], row_number, '镜号', integer=True))
            if original_index in used_indices:
                invalid(f'第 {row_number} 行镜号 {original_index} 重复，请填写唯一镜号')
            used_indices.add(original_index)
            duration = _positive_decimal(values['duration'], row_number, '时长')
            if not values['description']:
                invalid(f'第 {row_number} 行「画面描述」不能为空')
            index = len(scenes) + 1
            renumbered = renumbered or original_index != index
            duration_text = format(duration, 'f').rstrip('0').rstrip('.') if '.' in format(duration, 'f') else format(duration, 'f')
            scenes.append({
                **values, 'index': index, 'duration': duration_text + 's',
                'storyboard_prompt': values['final_prompt'], 'motion_prompt': '', 'resource_ids': [],
            })
        if not scenes:
            invalid('没有可导入的分镜，请在表头下方填写至少一行分镜内容')
        warnings = []
        if renumbered:
            warnings.append('已按表格行顺序将镜号重新编号为连续序号。')
        if extra_headers:
            warnings.append('已忽略模板以外的列：' + '、'.join(extra_headers[:10]))
        return {'scenes': scenes, 'count': len(scenes), 'sheet_name': sheet.title,
                'file_name': str(file_name).replace('\\', '/').rsplit('/', 1)[-1], 'warnings': warnings}
    except HTTPException:
        raise
    except Exception:
        invalid('Excel 文件已损坏或格式不受支持，请使用模板或重新另存为 .xlsx')
    finally:
        if workbook is not None:
            workbook.close()
